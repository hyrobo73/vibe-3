from __future__ import annotations

from pathlib import Path
from re import sub
from shutil import copyfileobj
from tempfile import NamedTemporaryFile
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import load_workbook

from app.core.config import settings
from app.db.connection import get_connection

router = APIRouter(prefix="/excel", tags=["excel"])

UPLOAD_DIR = settings.data_dir / "uploads"
EXPORT_DIR = settings.data_dir / "exports"
ALLOWED_EXTENSION = ".xlsx"
BLANK_GROUP_NAME = "__blank__"


def ensure_storage_dirs() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def require_xlsx(file: UploadFile) -> None:
    if Path(file.filename or "").suffix.lower() != ALLOWED_EXTENSION:
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")


def safe_filename(value: object, fallback: str = BLANK_GROUP_NAME) -> str:
    text = str(value).strip() if value is not None and str(value).strip() else fallback
    text = sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", text)
    text = sub(r"\s+", "_", text).strip(". ")
    return (text or fallback)[:80]


def save_upload(file: UploadFile) -> Path:
    require_xlsx(file)
    ensure_storage_dirs()
    suffix = Path(file.filename or "input.xlsx").suffix.lower()
    output_path = UPLOAD_DIR / f"{uuid4().hex}{suffix}"
    with output_path.open("wb") as output:
        copyfileobj(file.file, output)
    file.file.seek(0)
    return output_path


def read_headers(path: Path) -> list[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Unable to read the xlsx file") from exc

    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
    header_row = next(rows, None)
    workbook.close()

    if not header_row:
        raise HTTPException(status_code=400, detail="The first sheet does not contain a header row")

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    if not headers or any(header == "" for header in headers):
        raise HTTPException(status_code=400, detail="Header row cannot contain blank cells")
    if len(headers) != len(set(headers)):
        raise HTTPException(status_code=400, detail="Header row cannot contain duplicate column names")
    return headers


def read_dataframe(path: Path) -> pd.DataFrame:
    try:
        return pd.read_excel(path, sheet_name=0, engine="openpyxl", dtype=object)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Unable to read the xlsx file") from exc


def create_job(job_type: str, input_path: str | None = None) -> int:
    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO excel_jobs (job_type, status, input_path)
            VALUES (?, ?, ?)
            """,
            (job_type, "processing", input_path),
        )
    return int(cursor.lastrowid)


def update_job(job_id: int, status: str, output_path: str | None = None, error_message: str | None = None) -> None:
    with get_connection() as connection:
        connection.execute(
            """
            UPDATE excel_jobs
            SET status = ?, output_path = ?, error_message = ?
            WHERE id = ?
            """,
            (status, output_path, error_message, job_id),
        )


def get_job_or_404(job_id: int):
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, user_id, job_type, status, input_path, output_path, error_message, created_at
            FROM excel_jobs
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Excel job not found")
    return row


def job_to_response(row) -> dict[str, object]:
    output_path = row["output_path"]
    return {
        "id": row["id"],
        "job_type": row["job_type"],
        "status": row["status"],
        "error_message": row["error_message"],
        "created_at": row["created_at"],
        "download_url": f"/api/excel/jobs/{row['id']}/download" if output_path and row["status"] == "done" else None,
    }


def save_dataframe(df: pd.DataFrame, path: Path) -> None:
    df.to_excel(path, index=False, engine="openpyxl")


@router.post("/columns")
async def list_columns(file: UploadFile = File(...)) -> dict[str, object]:
    path = save_upload(file)
    headers = read_headers(path)
    return {"filename": file.filename, "columns": headers}


@router.post("/split", status_code=201)
async def split_excel(file: UploadFile = File(...), column_name: str = Form(...)) -> dict[str, object]:
    input_path = save_upload(file)
    job_id = create_job("split", str(input_path))

    try:
        headers = read_headers(input_path)
        if column_name not in headers:
            raise HTTPException(status_code=400, detail=f"Column '{column_name}' was not found")

        dataframe = read_dataframe(input_path)
        if dataframe.empty:
            raise HTTPException(status_code=400, detail="No data rows were found to split")

        dataframe = dataframe.reindex(columns=headers)
        group_values = dataframe[column_name].where(dataframe[column_name].notna(), BLANK_GROUP_NAME).map(safe_filename)
        groups: dict[str, pd.DataFrame] = {}
        for group_name in group_values.drop_duplicates():
            group_rows = dataframe.loc[group_values == group_name]
            if not group_rows.empty:
                groups[group_name] = group_rows

        if not groups:
            raise HTTPException(status_code=400, detail="No data rows were found to split")

        zip_path = EXPORT_DIR / f"excel_split_{job_id}_{uuid4().hex}.zip"
        with ZipFile(zip_path, "w", ZIP_DEFLATED) as zip_file:
            for group_name in sorted(groups):
                with NamedTemporaryFile(suffix=".xlsx", delete=False, dir=EXPORT_DIR) as temp_file:
                    temp_path = Path(temp_file.name)
                try:
                    save_dataframe(groups[group_name], temp_path)
                    zip_file.write(temp_path, arcname=f"{group_name}.xlsx")
                finally:
                    temp_path.unlink(missing_ok=True)

        update_job(job_id, "done", str(zip_path))
        return job_to_response(get_job_or_404(job_id))
    except HTTPException as exc:
        update_job(job_id, "failed", error_message=str(exc.detail))
        raise
    except Exception as exc:
        update_job(job_id, "failed", error_message="Failed to split the xlsx file")
        raise HTTPException(status_code=500, detail="Failed to split the xlsx file") from exc


@router.post("/merge", status_code=201)
async def merge_excels(files: list[UploadFile] = File(...)) -> dict[str, object]:
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="At least two .xlsx files are required for merge")

    input_paths = [save_upload(file) for file in files]
    job_id = create_job("merge", ";".join(str(path) for path in input_paths))

    try:
        base_headers = read_headers(input_paths[0])
        dataframes: list[pd.DataFrame] = []

        for index, path in enumerate(input_paths):
            headers = read_headers(path)
            if headers != base_headers:
                raise HTTPException(
                    status_code=400,
                    detail=f"Header mismatch in file {index + 1}. All files must have identical headers and order.",
                )

            dataframe = read_dataframe(path)
            if dataframe.empty:
                continue
            dataframes.append(dataframe.reindex(columns=base_headers))

        if not dataframes:
            raise HTTPException(status_code=400, detail="No data rows were found to merge")

        merged = pd.concat(dataframes, ignore_index=True)
        output_path = EXPORT_DIR / f"excel_merge_{job_id}_{uuid4().hex}.xlsx"
        save_dataframe(merged, output_path)
        update_job(job_id, "done", str(output_path))
        return job_to_response(get_job_or_404(job_id))
    except HTTPException as exc:
        update_job(job_id, "failed", error_message=str(exc.detail))
        raise
    except Exception as exc:
        update_job(job_id, "failed", error_message="Failed to merge the xlsx files")
        raise HTTPException(status_code=500, detail="Failed to merge the xlsx files") from exc


@router.get("/jobs/{job_id}")
def get_excel_job(job_id: int) -> dict[str, object]:
    return job_to_response(get_job_or_404(job_id))


@router.get("/jobs/{job_id}/download")
def download_excel_job(job_id: int) -> FileResponse:
    row = get_job_or_404(job_id)
    if row["status"] != "done" or not row["output_path"]:
        raise HTTPException(status_code=400, detail="The result file is not ready")

    output_path = Path(row["output_path"])
    if not output_path.exists():
        raise HTTPException(status_code=404, detail="The result file was not found")

    media_type = "application/zip" if output_path.suffix.lower() == ".zip" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return FileResponse(output_path, media_type=media_type, filename=output_path.name)
